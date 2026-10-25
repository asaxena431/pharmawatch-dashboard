"""Generate fake sample documents (no real data) into doc_redactor/samples/."""
import datetime as dt
import os

HERE = os.path.join(os.path.dirname(__file__), "samples")

LINES = [
    "ADVERSE EVENT CASE REPORT - CONFIDENTIAL",
    "Patient Name: John Q. Sample    MRN: 884213",
    "Date of Birth: 04/12/1961",
    "Address: 12 Oak Road, Rockville, MD 20850",
    "SSN 123-45-6789. Phone (301) 555-0199. Email john.sample@example.com",
    "Reporter: Dr. Maria Testa, reported on March 3, 2024.",
    "Product: Acme Pharma Zentrol 20 mg, LOT-AB1234 (Project Falcon).",
    "Narrative: patient developed a rash 2 days after the first dose on 2024-02-28.",
    "The event resolved after the drug was stopped. No other medications.",
]


def make_txt():
    with open(os.path.join(HERE, "sample_case_report.txt"), "w") as fh:
        fh.write("\n".join(LINES) + "\n")


def make_pdf():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for i, line in enumerate(LINES):
        page.insert_text((72, y), line, fontsize=14 if i == 0 else 11)
        y += 26 if i == 0 else 20
    doc.set_metadata({"author": "John Q. Sample", "title": "Case report 884213"})
    doc.save(os.path.join(HERE, "sample_case_report.pdf"))


def make_docx():
    import docx

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Acme Pharma - internal - Project Falcon"
    d.add_heading(LINES[0], 1)
    p = d.add_paragraph()
    p.add_run("Patient Name: ")
    p.add_run("John Q. ").bold = True
    p.add_run("Sample").bold = True
    for line in LINES[2:]:
        d.add_paragraph(line)
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for r, row in enumerate([("Field", "Value", "Notes"), ("SSN", "123-45-6789", "verified"),
                             ("Email", "john.sample@example.com", "call 301-555-0142")]):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.core_properties.author = "John Q. Sample"
    d.save(os.path.join(HERE, "sample_case_report.docx"))


def make_xlsx():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cases"
    ws.append(["Case ID", "Patient Name", "DOB", "SSN", "Phone", "Email", "Notes"])
    ws.append([1001, "John Q. Sample", dt.date(1961, 4, 12), "123-45-6789", "(301) 555-0199", "john.sample@example.com",
               "Acme Pharma lot LOT-AB1234; rash on 2024-02-28"])
    ws.append([1002, "Jane Roe", dt.date(1975, 9, 30), "987-65-4321", "240.555.0101", "jroe@example.org",
               "Project Falcon subject, headache 03/02/2024"])
    ws.append([1003, "Sam Placeholder", dt.date(1990, 1, 5), "111-22-3333", "202-555-0175", "sam.p@example.net",
               "No PHI in this note"])
    wb.properties.creator = "John Q. Sample"
    wb.save(os.path.join(HERE, "sample_case_list.xlsx"))


if __name__ == "__main__":
    os.makedirs(HERE, exist_ok=True)
    make_txt()
    make_pdf()
    make_docx()
    make_xlsx()
    print("samples written to", HERE)
