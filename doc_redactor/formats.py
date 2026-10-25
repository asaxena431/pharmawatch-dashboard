"""Apply a RuleSet to PDF, Word (.docx), Excel (.xlsx/.xlsm) and plain-text files."""
import datetime as dt
import io
import os
from dataclasses import dataclass, field

from .engine import RuleSet, find_matches, redact_text

TEXT_EXTS = {".txt", ".csv", ".tsv", ".md", ".log", ".json", ".xml", ".html"}
SUPPORTED_EXTS = sorted({".pdf", ".docx", ".xlsx", ".xlsm"} | TEXT_EXTS)
LEGACY_EXTS = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pdf"}

MIMETYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
}


class UnsupportedFile(ValueError):
    pass


@dataclass
class Hit:
    rule: str
    text: str
    location: str


@dataclass
class Result:
    data: bytes
    ext: str
    hits: list[Hit] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    preview_text: str = ""
    preview_images: list[bytes] = field(default_factory=list)

    @property
    def mimetype(self) -> str:
        return MIMETYPES.get(self.ext, "text/plain; charset=utf-8")


def output_name(filename: str) -> str:
    stem, ext = os.path.splitext(os.path.basename(filename))
    return f"{stem}_redacted{ext.lower()}"


def redact_file(filename: str, data: bytes, ruleset: RuleSet) -> Result:
    ext = os.path.splitext(filename)[1].lower()
    if ext in LEGACY_EXTS:
        raise UnsupportedFile(f"{ext} is an old binary format. Save it as {LEGACY_EXTS[ext]} and upload again.")
    if ext == ".pdf":
        return _pdf(data, ruleset)
    if ext == ".docx":
        return _docx(data, ruleset)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(data, ruleset, ext)
    if ext in TEXT_EXTS:
        return _text(data, ruleset, ext)
    raise UnsupportedFile(f"'{ext or filename}' is not supported. Supported: {', '.join(SUPPORTED_EXTS)}")


# ----------------------------------------------------------------- text -----
def _text(data: bytes, rs: RuleSet, ext: str) -> Result:
    try:
        text, enc = data.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError:
        text, enc = data.decode("latin-1"), "latin-1"
    hits = []
    out_lines = []
    for n, line in enumerate(text.splitlines(keepends=True), 1):
        new, ms = redact_text(line, rs)
        hits += [Hit(m.rule, m.text, f"line {n}") for m in ms]
        out_lines.append(new)
    out = "".join(out_lines)
    return Result(out.encode(enc), ext, hits, preview_text=out[:6000])


# ------------------------------------------------------------------ pdf -----
def _pdf(data: bytes, rs: RuleSet) -> Result:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise UnsupportedFile(f"could not open PDF: {exc}") from exc
    if doc.needs_pass:
        raise UnsupportedFile("PDF is password-protected; remove the password and upload again.")
    res = Result(b"", ".pdf")
    scanned = []
    for pno, page in enumerate(doc, 1):
        chars = []  # (char, rect or None, line_id)
        for bi, block in enumerate(page.get_text("rawdict")["blocks"]):
            if block.get("type") != 0:
                continue
            for li, line in enumerate(block["lines"]):
                for span in line["spans"]:
                    for ch in span["chars"]:
                        chars.append((ch["c"], pymupdf.Rect(ch["bbox"]), (bi, li)))
                chars.append(("\n", None, None))
        text = "".join(c for c, _, _ in chars)
        if not text.strip():
            if page.get_images():
                scanned.append(pno)
            continue
        for m in find_matches(text, rs):
            boxes: dict = {}
            for c, rect, line_id in chars[m.start:m.end]:
                if rect is None or c.isspace() or rect.is_empty:
                    continue
                boxes[line_id] = boxes[line_id] | rect if line_id in boxes else rect
            label = rs.replacement_for(m) if rs.pdf_label else None
            for rect in boxes.values():
                fontsize = max(4, min(9, rect.height * 0.75))
                page.add_redact_annot(rect, text=label, fontsize=fontsize, fill=(0, 0, 0), text_color=(1, 1, 1))
            res.hits.append(Hit(m.rule, m.text, f"page {pno}"))
        page.apply_redactions()
    if scanned:
        res.warnings.append(
            f"Page(s) {', '.join(map(str, scanned))} are scanned images with no text layer and were NOT redacted "
            "(OCR is not enabled in this version).")
    widgets = sum(1 for p in doc for _ in p.widgets())
    if widgets:
        res.warnings.append(f"{widgets} fillable form field(s) were not checked; flatten the form before redacting.")
    if rs.scrub_metadata:
        doc.set_metadata({})
        doc.del_xml_metadata()
    res.data = doc.tobytes(garbage=4, deflate=True, clean=True)
    out = pymupdf.open(stream=res.data, filetype="pdf")
    res.preview_images = [p.get_pixmap(matrix=pymupdf.Matrix(1.3, 1.3)).tobytes("png") for p in list(out)[:3]]
    res.preview_text = "\n".join(p.get_text() for p in out)[:6000]
    return res


# ----------------------------------------------------------------- docx -----
def _redact_runs(runs, rs: RuleSet, location: str, hits: list) -> None:
    texts = [r.text for r in runs]
    full = "".join(texts)
    matches = find_matches(full, rs)
    if not matches:
        return
    owner = [ri for ri, t in enumerate(texts) for _ in t]
    parts = [[] for _ in runs]
    i, mi = 0, 0
    while i < len(full):
        if mi < len(matches) and i == matches[mi].start:
            m = matches[mi]
            parts[owner[i]].append(rs.replacement_for(m))
            hits.append(Hit(m.rule, m.text, location))
            i, mi = m.end, mi + 1
        else:
            parts[owner[i]].append(full[i])
            i += 1
    for run, p in zip(runs, parts):
        if run.text != "".join(p):
            run.text = "".join(p)


def _paragraph_runs(paragraph):
    from docx.text.hyperlink import Hyperlink

    runs = []
    for item in paragraph.iter_inner_content():
        runs.extend(item.runs if isinstance(item, Hyperlink) else [item])
    return runs


def _docx(data: bytes, rs: RuleSet) -> Result:
    import docx
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedFile(f"could not open Word file: {exc}") from exc
    res = Result(b"", ".docx")
    seen = set()

    def walk(root_el, parent, label):
        for n, p_el in enumerate(root_el.iter(qn("w:p")), 1):
            if id(p_el) in seen:
                continue
            seen.add(id(p_el))
            _redact_runs(_paragraph_runs(Paragraph(p_el, parent)), rs, f"{label} paragraph {n}", res.hits)

    walk(document.element.body, document._body, "body")
    for si, section in enumerate(document.sections, 1):
        for kind in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
            part = getattr(section, kind)
            if not part.is_linked_to_previous:
                walk(part._element, part, f"section {si} {kind.replace('_', ' ')}")
    if rs.scrub_metadata:
        cp = document.core_properties
        for attr in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
            setattr(cp, attr, "")
    buf = io.BytesIO()
    document.save(buf)
    res.data = buf.getvalue()
    out = docx.Document(io.BytesIO(res.data))
    lines = [p.text for p in out.paragraphs]
    for t in out.tables:
        lines += [" | ".join(c.text for c in row.cells) for row in t.rows]
    res.preview_text = "\n".join(lines)[:6000]
    return res


# ----------------------------------------------------------------- xlsx -----
def _xlsx(data: bytes, rs: RuleSet, ext: str) -> Result:
    import openpyxl
    from openpyxl.comments import Comment

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), keep_vba=ext == ".xlsm")
    except Exception as exc:
        raise UnsupportedFile(f"could not open Excel file: {exc}") from exc
    res = Result(b"", ext)
    formulas = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                loc = f"{ws.title}!{cell.coordinate}"
                v = cell.value
                if isinstance(v, str):
                    if v.startswith("="):
                        formulas += 1
                        continue
                    new, ms = redact_text(v, rs)
                    if ms:
                        cell.value = new
                        res.hits += [Hit(m.rule, m.text, loc) for m in ms]
                elif isinstance(v, (dt.datetime, dt.date)):
                    as_text = v.strftime("%m/%d/%Y")
                    ms = find_matches(as_text, rs)
                    if ms:
                        cell.value = rs.replacement_for(ms[0])
                        cell.number_format = "General"
                        res.hits.append(Hit(ms[0].rule, as_text, loc))
                if cell.comment is not None:
                    new, ms = redact_text(cell.comment.text, rs)
                    if ms:
                        cell.comment = Comment(new, "")
                        res.hits += [Hit(m.rule, m.text, f"{loc} comment") for m in ms]
    if formulas:
        res.warnings.append(f"{formulas} formula cell(s) were left unchanged.")
    if rs.scrub_metadata:
        p = wb.properties
        p.creator = p.lastModifiedBy = p.title = p.subject = p.keywords = p.description = p.category = None
    buf = io.BytesIO()
    wb.save(buf)
    res.data = buf.getvalue()
    out = openpyxl.load_workbook(io.BytesIO(res.data), read_only=True)
    lines = []
    for ws in out.worksheets:
        lines.append(f"--- sheet: {ws.title} ---")
        for r in ws.iter_rows(max_row=40, values_only=True):
            lines.append("\t".join("" if v is None else str(v) for v in r))
    res.preview_text = "\n".join(lines)[:6000]
    return res
