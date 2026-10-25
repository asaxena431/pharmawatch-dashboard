import io
import os

import openpyxl
import pymupdf
import pytest
from docx import Document

from doc_redactor.engine import RulesError, load_rules, load_rules_file, redact_text
from doc_redactor.formats import UnsupportedFile, redact_file

RULES = load_rules_file(os.path.join(os.path.dirname(__file__), "..", "config", "rules.yaml"))


@pytest.mark.parametrize("text,expected", [
    ("Patient Name: John Q. Doe    MRN: 884213", "Patient Name: [PHI]    MRN: [PHI]"),
    ("Reporter: Dr. Maria Testa", "Reporter: [PHI]"),
    ("Name: John Doe. He was admitted.", "Name: [PHI]. He was admitted."),
    ("Patient Name: Jane Roe; Address: 12 Main St.", "Patient Name: [PHI]; Address: [PHI]."),
    ("Date of Birth: 04/12/1961\tSex: F", "Date of Birth: [DATE]\tSex: F"),
    ("SSN 123-45-6789, mail a.b@example.com, call (301) 555-0101", "SSN [SSN], mail [EMAIL], call [PHONE]"),
    ("Made by ACME pharma under project falcon", "Made by [CONFIDENTIAL] under [CONFIDENTIAL]"),
    ("Lot LOT-AB1234 shipped", "Lot LOT-[LOT] shipped"),
])
def test_default_rules(text, expected):
    assert redact_text(text, RULES)[0] == expected


def test_disabled_rule_and_mask():
    rs = load_rules("settings: {mask_char: '#'}\nrules:\n- {name: ssn, type: builtin}\n- {name: email, type: builtin, enabled: false}")
    assert redact_text("123-45-6789 x@y.com", rs)[0] == "########### x@y.com"


@pytest.mark.parametrize("bad", [
    "rules: [{name: x, type: builtin}]",
    "rules: [{name: x, type: regex, pattern: '('}]",
    "rules: [{name: x, type: keywords, values: []}]",
    "rules: [{name: x, type: nope}]",
    "rules: [{name: a, type: builtin, pattern_name: ssn}, {name: a, type: builtin, pattern_name: email}]",
    "rules: : :",
])
def test_bad_rules(bad):
    with pytest.raises(RulesError):
        load_rules(bad)


def test_pdf_text_removed():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Patient SSN 123-45-6789 at Acme Pharma")
    res = redact_file("x.pdf", doc.tobytes(), RULES)
    text = pymupdf.open(stream=res.data, filetype="pdf")[0].get_text()
    assert "123-45-6789" not in text and "Acme" not in text
    assert {h.rule for h in res.hits} == {"ssn", "confidential_terms"}


def test_docx_split_runs():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("SSN 123-")
    p.add_run("45-6789 end")
    buf = io.BytesIO()
    doc.save(buf)
    res = redact_file("x.docx", buf.getvalue(), RULES)
    out = Document(io.BytesIO(res.data))
    assert out.paragraphs[0].text == "SSN [SSN] end"


def test_xlsx_cells_and_formulas():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "email me at a@b.com"
    ws["A2"] = "=1+1"
    buf = io.BytesIO()
    wb.save(buf)
    res = redact_file("x.xlsx", buf.getvalue(), RULES)
    out = openpyxl.load_workbook(io.BytesIO(res.data)).active
    assert out["A1"].value == "email me at [EMAIL]"
    assert out["A2"].value == "=1+1"


def test_legacy_rejected():
    with pytest.raises(UnsupportedFile):
        redact_file("old.doc", b"x", RULES)
