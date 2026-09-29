"""Benchmark the redactor against narratives that a human editor already redacted.

Input: an .xlsx / .csv with (at least) an original-text column and a human-redacted
column (defaults match the FDA export: COMPANY_NARRATIVE / REDACTED_NARRATIVE).
For every row the original is redacted, compared with the human version and the
difference is explained with a reason category so a reviewer can see *why* the
automation and the editor disagree.

    python -m mdr_redaction.compare redacted_narratives.xlsx -o compare_out.xlsx
"""
from __future__ import annotations

import argparse
import csv
import difflib
import re
import sys
from collections import Counter
from dataclasses import dataclass, asdict
from pathlib import Path

from . import rules as R
from .redactor import redact_section

_TOKEN = re.compile(r"\(b\)\s*\)?\((\d)\)", re.I)          # (b)(6), (B)(4), (b))(6) typo
_WS = re.compile(r"\s+")
_PUNCT = re.compile(r"[.,;:]+(?=\s|$)")

_DATE_LIKE = re.compile(
    rf"\b(?:\d{{1,2}}[/.-]\d{{1,2}}[/.-]\d{{2,4}}|\d{{4}}-\d{{1,2}}-\d{{1,2}}|\d{{1,2}}-{R._MONTHS}-\d{{2,4}}"
    rf"|{R._MONTHS}\.?\s+\d{{1,2}},?\s+\d{{4}}|\d{{1,2}}\s+{R._MONTHS}\.?,?\s+\d{{4}})\b", re.I)


def _norm(s: str | None) -> str:
    s = _TOKEN.sub(lambda m: f"(b)({m.group(1)})", s or "")
    s = re.sub(r"\(b\)\((\d)\)(?=\S)", r"(b)(\1) ", s)            # "(b)(6)2026" -> "(b)(6) 2026"
    s = re.sub(r"(\d{4})\s*-?\s*\(b\)\((\d)\)", r"(b)(\2) \1", s)  # "2026-(b)(6)" -> "(b)(6) 2026"
    s = _PUNCT.sub("", s)
    return _WS.sub(" ", s).strip().lower()


def _tokens(s: str) -> list[str]:
    return _norm(s).split(" ")


@dataclass
class RowResult:
    record_id: str
    status: str            # MATCH | FORMAT_ONLY | AUTO_MISSED | AUTO_EXTRA | BOTH | HUMAN_INCONSISTENT
    reason: str
    auto_only: str         # text the automation redacted but the editor kept
    human_only: str        # text the editor redacted but the automation kept
    original: str
    human_redacted: str
    auto_redacted: str
    findings: str


def _redacted_spans(original: str, redacted: str) -> tuple[list[tuple[str, str]], list[str]]:
    """Align the original with a redacted version.

    Returns (redactions, edits): `redactions` are (original words, replacement) pairs where the
    replacement contains a (b)(n) token; `edits` are other wording changes (typo fixes,
    punctuation, added words) that are not redactions.
    """
    to, tr = _tokens(original), _tokens(redacted)
    sm = difflib.SequenceMatcher(None, to, tr, autojunk=False)
    redactions, edits = [], []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        src, dst = " ".join(to[i1:i2]), " ".join(tr[j1:j2])
        if "(b)(" in dst:
            redactions.append((src, dst))
        elif src.strip("-*") or dst.strip("-*"):
            edits.append(f"'{src}' -> '{dst}'")
    return redactions, edits


def _classify_missed(span: str) -> str:
    if _DATE_LIKE.search(span):
        return f"date format not recognised by automation: '{span}'"
    if re.fullmatch(rf"(?:{R._MONTHS}\.?\s*)?\d{{0,2}}\s*(?:of\s+)?(?:19|20)?\d{{0,4}}", span, re.I) and re.search(r"[a-z]", span, re.I):
        return f"partial date (month/day without full date) redacted by editor: '{span}'"
    if re.fullmatch(r"(19|20)\d{2}", span):
        return f"editor removed the year '{span}' (SOP Appendix 6 keeps the year unless patient >89)"
    if re.search(r"\d", span):
        return f"editor redacted identifier/number not covered by a rule: '{span}'"
    return f"editor redacted free text (name/facility/location?) that no rule detects: '{span}'"


def _explain(original: str, human: str, auto: str, findings) -> tuple[str, str, str, str]:
    if human.strip() == auto.strip():
        return "MATCH", "identical", "", ""
    if _norm(human) == _norm(auto):
        return "FORMAT_ONLY", "same redactions; differs only in (b)(6) casing/spacing/punctuation", "", ""

    h_red, h_edits = _redacted_spans(original, human)
    a_red, _ = _redacted_spans(original, auto)
    h_src = Counter(s for s, _ in h_red)
    a_src = Counter(s for s, _ in a_red)
    human_only = list((h_src - a_src).elements())     # editor redacted, automation kept
    auto_only = list((a_src - h_src).elements())      # automation redacted, editor kept

    reasons = []
    # same source span, different replacement (e.g. year kept vs dropped)
    def _core(d: str) -> tuple:  # only the tokens and years matter, not filler words like "on"
        return tuple(re.findall(r"\(b\)\(\d\)|(?:19|20)\d{2}", d))

    for src in (h_src & a_src):
        h_rep = {d for s, d in h_red if s == src}
        a_rep = {d for s, d in a_red if s == src}
        if {_core(d) for d in h_rep} == {_core(d) for d in a_rep}:
            if h_rep != a_rep:
                h_edits.append(f"'{src}': editor wrote {sorted(h_rep)} vs {sorted(a_rep)}")
        else:
            if any(re.search(r"(19|20)\d{2}", d) for d in a_rep) and not any(re.search(r"(19|20)\d{2}", d) for d in h_rep):
                reasons.append(f"editor dropped the year for '{src}' (SOP Appendix 6 keeps the year unless patient >89)")
            else:
                reasons.append(f"'{src}': editor wrote {sorted(h_rep)}, automation wrote {sorted(a_rep)}")

    for span in human_only:
        reasons.append(_classify_missed(span))
    for span in auto_only:
        kinds = sorted({f.rule for f in findings if _norm(f.original) and (_norm(f.original) in span or span in _norm(f.original))})
        if _DATE_LIKE.search(span) or kinds == ["date"]:
            reasons.append(f"automation redacted date '{span}' but editor left it "
                           "(editor inconsistency: same date type is redacted elsewhere)")
        elif kinds:
            reasons.append(f"automation rule {kinds} fired on '{span}' but editor kept it (probable false positive)")
        else:
            reasons.append(f"automation redacted '{span}' but editor kept it")

    if human_only and auto_only:
        status = "BOTH"
    elif human_only:
        status = "AUTO_MISSED"
    elif auto_only:
        status = "AUTO_EXTRA"
    elif reasons:
        status = "HUMAN_INCONSISTENT"
    else:
        status = "FORMAT_ONLY"
        reasons.append("same redactions; editor only reworded/re-punctuated the text")

    if h_edits:
        reasons.append("editor wording edits (not redactions): " + ", ".join(h_edits[:5]))

    return status, "; ".join(reasons), " | ".join(auto_only), " | ".join(human_only)


def load_rows(path: Path, orig_col: str, red_col: str, id_col: str) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        ws = openpyxl.load_workbook(path, read_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(it)]
        rows = [dict(zip(header, r)) for r in it]
    else:
        with open(path, newline="", encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
    for c in (orig_col, red_col):
        if rows and c not in rows[0]:
            sys.exit(f"column {c!r} not found; have {list(rows[0])}")
    return [r for r in rows if r.get(orig_col)]


def compare(rows: list[dict], orig_col: str, red_col: str, id_col: str, section: str = "B5") -> list[RowResult]:
    out = []
    for r in rows:
        orig, human = str(r[orig_col]), str(r.get(red_col) or "")
        res = redact_section(section, orig)
        status, reason, auto_only, human_only = _explain(orig, human, res.redacted, res.findings)
        out.append(RowResult(str(r.get(id_col, "")), status, reason, auto_only, human_only,
                             orig, human, res.redacted,
                             "; ".join(f"{f.rule}:{f.original!r}->{f.replacement!r}" for f in res.findings)))
    return out


def summary(results: list[RowResult]) -> str:
    c = Counter(r.status for r in results)
    n = len(results)
    agree = c["MATCH"] + c["FORMAT_ONLY"]
    lines = [f"rows: {n}", f"agree with editor (exact or format-only): {agree} ({agree / n:.0%})"]
    for k in ("MATCH", "FORMAT_ONLY", "AUTO_MISSED", "AUTO_EXTRA", "BOTH", "HUMAN_INCONSISTENT"):
        lines.append(f"  {k:<19}{c[k]}")
    return "\n".join(lines)


def write_xlsx(results: list[RowResult], path: Path) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "comparison"
    cols = list(RowResult.__dataclass_fields__)
    ws.append(cols)
    for r in results:
        ws.append([asdict(r)[c] for c in cols])
    ws2 = wb.create_sheet("summary")
    for line in summary(results).splitlines():
        ws2.append([line.strip()])
    wb.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", type=Path)
    ap.add_argument("-o", "--output", type=Path, default=Path("compare_out.xlsx"))
    ap.add_argument("--orig-col", default="COMPANY_NARRATIVE")
    ap.add_argument("--redacted-col", default="REDACTED_NARRATIVE")
    ap.add_argument("--id-col", default="RECORD_ID")
    ap.add_argument("--section", default="B5")
    a = ap.parse_args(argv)
    results = compare(load_rows(a.input, a.orig_col, a.redacted_col, a.id_col), a.orig_col, a.redacted_col, a.id_col, a.section)
    if a.output.suffix.lower() == ".csv":
        with open(a.output, "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(RowResult.__dataclass_fields__))
            w.writeheader()
            w.writerows(asdict(r) for r in results)
    else:
        write_xlsx(results, a.output)
    print(summary(results))
    print(f"written {a.output}")


if __name__ == "__main__":
    main()
