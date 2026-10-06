import unittest
from pathlib import Path

from mdr_redaction.compare import _explain, compare
from mdr_redaction.redactor import redact_section


class DateFormatsFromRealExport(unittest.TestCase):
    def test_dd_mon_yyyy(self):
        self.assertEqual(redact_section("B5", "low glucose by 05-Nov-2025.").redacted, "low glucose by (B)(6) 2025.")

    def test_typo_three_digit_year(self):
        self.assertEqual(redact_section("B5", "ON 04-DEC-2018 AND/OR 14-DEC-021(COULD BE ERROR)").redacted,
                         "ON (B)(6) 2018 AND/OR (B)(6) 2021(COULD BE ERROR)")

    def test_dd_slash_mon_slash_yyyy(self):
        self.assertEqual(redact_section("B5", "ON 05/DEC/2025, THE PATIENT").redacted, "ON (B)(6) 2025, THE PATIENT")

    def test_month_of_year(self):
        self.assertEqual(redact_section("B5", "surgery in August of 2019.").redacted, "surgery in (B)(6) of 2019.")

    def test_month_day_without_year(self):
        self.assertEqual(redact_section("B5", "PASSED AWAY ON MARCH 25 AFTER").redacted, "PASSED AWAY ON (B)(6) AFTER")


class FalsePositivesFromRealExport(unittest.TestCase):
    def test_patient_is_not_initials(self):
        t = "THE PATIENT IS MAKING A CLAIM"
        self.assertEqual(redact_section("B5", t).redacted, t)

    def test_pt_abbreviation_not_initials(self):
        t = "the patient (pt) reported pain"
        self.assertEqual(redact_section("B5", t).redacted, t)

    def test_sneezed_not_serial(self):
        t = "PATIENT SNEEZED AND WAS ADMITTED"
        self.assertEqual(redact_section("B5", t).redacted, t)

    def test_serialized_not_serial(self):
        t = "the device is serialized and tracked"
        self.assertEqual(redact_section("B5", t).redacted, t)

    def test_pain_at_rest_not_address(self):
        t = "score 93.1 pain at rest"
        self.assertEqual(redact_section("B5", t).redacted, t)

    def test_uppercase_mr_name(self):
        self.assertEqual(redact_section("B5", "MR. RAFAEL BENITEZ REPORTED").redacted, "(B)(6) REPORTED")


class Comparison(unittest.TestCase):
    def test_format_only(self):
        orig = "implant inserted 2026-02-09 in FDI 46."
        human = "implant inserted on (b)(6)2026 in FDI 46."
        auto = redact_section("B5", orig).redacted
        status, reason, _, _, _ = _explain(orig, human, auto, [])
        self.assertEqual(status, "FORMAT_ONLY")

    def test_auto_missed_free_text(self):
        orig = "seen at ZORBA WELLNESS STUDIO on 05/05/2026"
        human = "seen at (b)(6) on (b)(6) 2026"
        res = redact_section("B5", orig)
        status, reason, ref, auto_only, human_only = _explain(orig, human, res.redacted, res.findings)
        self.assertEqual(status, "AUTO_MISSED")
        self.assertIn("zorba wellness studio", human_only)
        self.assertIn("Appendix 6", ref)

    def test_auto_extra_reports_rule(self):
        orig = "event on 2026-06-10 verified"
        human = "event on 2026-06-10 verified"      # editor left the date
        res = redact_section("B5", orig)
        status, reason, ref, auto_only, _ = _explain(orig, human, res.redacted, res.findings)
        self.assertEqual(status, "AUTO_EXTRA")
        self.assertIn("editor inconsistency", reason)
        self.assertIn("Appendix 6", ref)

    def test_compare_rows(self):
        rows = [{"RECORD_ID": "1", "COMPANY_NARRATIVE": "on 06/25/2015 pt fell", "REDACTED_NARRATIVE": "on (b)(6) 2015 pt fell"}]
        out = compare(rows, "COMPANY_NARRATIVE", "REDACTED_NARRATIVE", "RECORD_ID")
        self.assertEqual(out[0].status, "FORMAT_ONLY")


if __name__ == "__main__":
    unittest.main()


class SecondExport1000Rows(unittest.TestCase):
    def test_compact_ddmonyyyy(self):
        self.assertEqual(redact_section("B5", "seen on 22Jun2026 and 27AUG2026.").redacted, "seen on (B)(6) 2026 and (B)(6) 2026.")

    def test_device_identifier_and_ref_numbers(self):
        r = redact_section("H11", "Device Identifier: 00650862130287. Manufacturer's Ref. #: 1555899. "
                                  "FILED UNDER MFR NUMBER (1822565). report 1 of 2 for PC-002241284").redacted
        self.assertNotIn("00650862130287", r)
        self.assertNotIn("1555899", r)
        self.assertNotIn("1822565", r)
        self.assertNotIn("002241284", r)

    def test_false_positives_from_second_export(self):
        for text in ("The distributor reported foreign matter inside the package.",
                     "A US distributor contacted ZOLL to report that a patient developed a rash.",
                     "the serial number associated with this complaint",
                     "a transmitter battery issue occurred",
                     "Account noticed the right front wheel",
                     "PURSUANT TO 21 CFR PART 803",
                     "Straight - Model Number 8028.",
                     "Another Cardiology MD was called.",
                     "the patient Insulin Doesn't Exit",
                     "(Manufacturer Representative, Healthcare Provider) regarding",
                     "Per Op Notes the incision"):
            self.assertEqual(redact_section("B5", text).redacted, text, text)

    def test_month_year_and_emergency_center(self):
        self.assertEqual(redact_section("B5", "In September 2025, the patient").redacted, "In (B)(6) 2025, the patient")
        self.assertEqual(redact_section("B5", "transported to Tidal Health Emergency Center in Berlin, MD and released").redacted,
                         "transported to (B)(6) Emergency Center and released")

    def test_bare_month_and_month_year(self):
        self.assertEqual(redact_section("B5", "chair received back in March has broken").redacted,
                         "chair received back in (B)(6) has broken")
        self.assertEqual(redact_section("B5", "In September 2025, the patient").redacted, "In (B)(6) 2025, the patient")
        self.assertEqual(redact_section("B5", "it may be in place").redacted, "it may be in place")

    def test_editor_over_redaction_is_not_auto_missed(self):
        from mdr_redaction.compare import _explain
        orig = "Drive Medical was notified. Results of 50 mg/dL were seen."
        human = "(b)(6) was notified. Results of (b)(6) were seen."
        status, reason, ref, *_ = _explain(orig, human, orig, [])
        self.assertEqual(status, "HUMAN_INCONSISTENT")
        self.assertIn("over-redaction", reason)
        self.assertIn("[not_pii]", ref)


class WhitespaceOnlyTests(unittest.TestCase):
    def test_double_space_is_match(self):
        from mdr_redaction.compare import _explain
        orig = "Stent snapped.  Another stent was placed."
        status, reason, *_ = _explain(orig, orig, "Stent snapped. Another stent was placed.", [])
        self.assertEqual(status, "MATCH")


class LargeExportFormats(unittest.TestCase):
    def test_new_date_and_id_formats(self):
        out = redact_section("B5", "ON 2026-FEB-04 AND 21-APR2026 (APR-2026, 2026/06/03), SEEN ON 7/9. CRD_1032 REGISTRY, "
                             "CLINICAL ID 000078-013, TRIAL (BPV18002). SUBMITTED UNDER 3007963827. "
                             "AT MIRAMAR FAMILY DENTAL AND BARNES-JEWISH HOSPITAL.").redacted
        for kept in ("2026-FEB", "APR2026", "APR-2026", "2026/06", "7/9", "CRD_1032", "000078", "BPV18002", "3007963827",
                     "MIRAMAR", "BARNES"):
            self.assertNotIn(kept, out, out)
        self.assertIn("(B)(6) 2026", out)
        self.assertEqual(redact_section("B5", "RATIO 7/10 OF PATIENTS, BP 120/80, TAKEN TO THE HOSPITAL.").redacted,
                         "RATIO 7/10 OF PATIENTS, BP 120/80, TAKEN TO THE HOSPITAL.")

    def test_caps_false_positives(self):
        txt = ("A US DISTRIBUTOR RETURNED A MONITOR. A US DISTRIBUTOR REPORTED THAT A PATIENT (PT) HAD GRADE 4 MR. "
               "MR WAS REDUCED TO GRADE 1. THE SUPPLIER IS Acme Plastics.")
        out = redact_section("B5", txt).redacted
        self.assertIn("DISTRIBUTOR RETURNED", out)
        self.assertIn("DISTRIBUTOR REPORTED THAT", out)
        self.assertIn("(PT)", out)
        self.assertIn("MR WAS REDUCED", out)
        self.assertIn("SUPPLIER IS (B)(4)", out)


class XlsxColouring(unittest.TestCase):
    def test_rich_text_colours(self):
        import tempfile
        import openpyxl
        from openpyxl.cell.rich_text import CellRichText
        from mdr_redaction.compare import RowResult, write_xlsx, RED, BLUE
        r = RowResult("1", "BOTH", "r", "ref", "Model 8028", "Acme Clinic",
                      "Seen at Acme Clinic with Model 8028 on 1/2/2026", "Seen at (b)(6) with Model 8028 on (b)(6) 2026",
                      "Seen at Acme Clinic with Model (B)(4) on (B)(6) 2026", "")
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "o.xlsx"
            write_xlsx([r], p)
            ws = openpyxl.load_workbook(p, rich_text=True)["comparison"]
            cols = [c.value for c in ws[1]]
            orig = ws.cell(row=2, column=cols.index("original") + 1).value
            self.assertIsInstance(orig, CellRichText)
            colours = {str(b.text): b.font.color.rgb[-6:] for b in orig if hasattr(b, "font")}
            self.assertEqual(colours["Acme Clinic"], RED)
            self.assertEqual(colours["Model 8028"], BLUE)
            self.assertEqual(str(orig), r.original)
